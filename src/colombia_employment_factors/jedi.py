"""Read saved JEDI caches without recalculating Excel or executing VBA."""

from __future__ import annotations

import hashlib
import math
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pandas as pd

SOURCE_COLUMNS = ["Source", "Technology", "Factor_Type", "Job_Type", "Unit", "Year", "Value", "Value_Numeric"]
JOBS = ("Direct", "Indirect", "Induced")
PHASES = ("Construction&Manufacturing", "O&M")
I_JEDI_TECHNOLOGIES = {
    "Wind": "Onshore wind", "Solar-PV": "Utility-scale solar PV",
    "Biopower": "Biomass power", "Geothermal": "Geothermal",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _formula(cell):
    value = cell.value
    return {"text": value.text, "array_range": value.ref} if hasattr(value, "text") else value


def _number(value, label: str, *, positive: bool = False) -> float:
    if isinstance(value, bool):
        raise ValueError(f"Invalid numeric JEDI cache: {label}")
    try:
        result = float(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"Missing/invalid JEDI cache: {label}") from exc
    if not math.isfinite(result) or result < 0 or (positive and result == 0):
        raise ValueError(f"Invalid numeric JEDI cache: {label}")
    return result


def _record(technology: str, source: str, phase: str, category: str, value: float) -> dict:
    return dict(zip(SOURCE_COLUMNS, [source, technology, phase, category,
                "job/MW" if phase == "O&M" else "job-yr/MW", 2024, value, value]))


def _metadata(path: Path, workbook, root: Path) -> dict:
    with ZipFile(path) as archive:
        names = archive.namelist()
    return {
        "path": path.relative_to(root).as_posix(), "sha256": sha256(path),
        "hidden_sheets": [s.title for s in workbook if s.sheet_state != "visible"],
        "named_ranges": sorted(workbook.defined_names),
        "external_links": [n for n in names if n.startswith("xl/externalLinks/externalLink")],
        "contains_vba": "xl/vbaProject.bin" in names,
        "calculation_id": workbook.calculation.calcId if workbook.calculation else None,
        "cache_validation": "Saved caches only; input identities and capacity/result totals checked; not a full Excel recalculation",
    }


def extract_i_jedi(path: Path, root: Path) -> tuple[pd.DataFrame, dict]:
    values = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        manifest = {**_metadata(path, formulas, root), "runs": []}
        rows = []
        for prefix, technology in I_JEDI_TECHNOLOGIES.items():
            inputs, results = values[prefix + "-Inputs"], values[prefix + "-Results"]
            year_cell = {"Biopower": "B9", "Geothermal": "B10"}.get(prefix, "B6")
            if inputs["B4"].value != "Colombia" or inputs[year_cell].value != 2024:
                raise ValueError(f"Expected Colombia / 2024 money in {inputs.title}")
            capacity = _number(inputs["B5"].value, inputs.title + "!B5", positive=True)
            run = {"technology": technology, "country": "Colombia", "money_year": 2024,
                   "capacity_mw": capacity, "capacity_cell": inputs.title + "!B5", "results": []}
            for phase, first_row, total_row in [(PHASES[0], 14, 17), (PHASES[1], 21, 24)]:
                counts = [_number(results[f"C{first_row+i}"].value, f"{results.title}!C{first_row+i}") for i in range(3)]
                total = _number(results[f"C{total_row}"].value, f"{results.title}!C{total_row}")
                if not math.isclose(sum(counts), total, rel_tol=1e-9, abs_tol=1e-8):
                    raise ValueError(f"Inconsistent saved result totals: {results.title}/{phase}")
                for i, (category, count) in enumerate(zip(JOBS, counts)):
                    row = _record(technology, "I-JEDI", phase, category, count / capacity)
                    rows.append(row)
                    cell = f"C{first_row+i}"
                    run["results"].append({"phase": phase, "category": category,
                        "cell": results.title + "!" + cell, "formula": _formula(formulas[results.title][cell]),
                        "cached_jobs": count, "factor_per_mw": row["Value_Numeric"]})
            manifest["runs"].append(run)
        return pd.DataFrame(rows, columns=SOURCE_COLUMNS), manifest
    finally:
        values.close()
        formulas.close()


def extract_us_jedi(path: Path, root: Path, configuration: str, *, expected_capacity_mw: float) -> tuple[pd.DataFrame, dict]:
    values = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        inputs, results = values["ProjectData"], values["SummaryResults"]
        version = values["About JEDI"]["A6"].value
        if "PV05.20.21" not in str(version):
            raise ValueError("Unsupported US-JEDI release; review the extraction cell map")
        settings = {c: inputs[c].value for c in ["B12", "B14", "B15", "B16", "B17", "B18", "B19", "B20", "B22", "B24", "B26"]}
        if settings["B12"] != "UNITED STATES" or settings["B14"] != 2024 or settings["B15"] != configuration:
            raise ValueError(f"Expected US / 2024 / {configuration} configuration: {path.name}")
        capacity_kw = _number(settings["B20"], "ProjectData!B20", positive=True)
        calculated = _number(settings["B18"], "ProjectData!B18", positive=True) * _number(settings["B19"], "ProjectData!B19", positive=True)
        if not math.isclose(capacity_kw, calculated, rel_tol=1e-12):
            raise ValueError("Stale US-JEDI capacity cache: total differs from system size times system count")
        capacity_mw = capacity_kw / 1000
        expected_capacity_mw = _number(expected_capacity_mw, "approved US-JEDI capacity", positive=True)
        if not math.isclose(capacity_mw, expected_capacity_mw, rel_tol=1e-12):
            raise ValueError(f"Expected the approved {expected_capacity_mw:g} MW US-JEDI {configuration} run; review configuration rather than silently rescaling")
        for cell, expected in [("C3", settings["B14"]), ("C6", capacity_kw), ("C7", configuration), ("C12", settings["B24"])]:
            if results[cell].value != expected:
                raise ValueError(f"Inconsistent US-JEDI summary cache: {cell}")
        technology = "Rooftop solar PV" if configuration == "Residential" else "Utility-scale solar PV"
        rows, references = [], []
        for phase, cells, total_cell in [(PHASES[0], ["B30", "B38", "B39"], "B40"), (PHASES[1], ["B46", "B47", "B48"], "B49")]:
            counts = [_number(results[c].value, "SummaryResults!" + c) for c in cells]
            total = _number(results[total_cell].value, "SummaryResults!" + total_cell)
            if not math.isclose(sum(counts), total, rel_tol=1e-9, abs_tol=1e-8):
                raise ValueError(f"Inconsistent US-JEDI saved result totals: {phase}")
            for cell, category, count in zip(cells, JOBS, counts):
                rows.append(_record(technology, "JEDI-US", phase, category, count / capacity_mw))
                references.append({"phase": phase, "category": category, "cell": "SummaryResults!" + cell,
                    "formula": _formula(formulas["SummaryResults"][cell]), "cached_jobs": count,
                    "factor_per_mw": count / capacity_mw})
        return pd.DataFrame(rows, columns=SOURCE_COLUMNS), {
            **_metadata(path, formulas, root), "version": version, "settings": settings,
            "capacity_mw": capacity_mw, "expected_capacity_mw": expected_capacity_mw,
            "capacity_basis": "MW DC; ProjectData!B18 (kW DC/system) * B19 (systems) / 1000",
            "normalisation": "Saved total jobs divided once by actual total capacity_mw before forming ratios",
            "monetary_evidence": {
                "base_cost_per_kw_dc": settings["B22"],
                "base_cost_formula": _formula(formulas["ProjectData"]["B22"]),
                "dollar_year_cell": "ProjectData!B24",
                "multiplier_reference_year": values["Multipliers"]["D2"].value,
                "deflator_to_multiplier_year": values["Deflators"]["F7"].value,
                "deflator_formula": _formula(formulas["Deflators"]["F7"]),
                "deflator_note": values["Deflators"]["C45"].value,
                "interpretation": "Shared saved dollar-year convention for relative transfer; not verification of default cost vintage or 2024 labour productivity",
            },
            "results": references,
            "manufacturing_subtotal": _number(results["B32"].value, "SummaryResults!B32"),
            "boundary_map": {"on-site labour and services": "Direct", "supply chain": "Indirect", "induced": "Induced"},
        }
    finally:
        values.close()
        formulas.close()


def extract_jedi_sources(root: Path, policy: dict) -> tuple[pd.DataFrame, dict]:
    files = policy["jedi_workbooks"]
    frames, manifest = [], {}
    frame, manifest["i_jedi"] = extract_i_jedi(root / files["i_jedi"], root)
    frames.append(frame)
    capacities = policy["us_jedi_expected_capacity_mw"]
    frame, manifest["us_residential"] = extract_us_jedi(root / files["us_residential"], root, "Residential", expected_capacity_mw=capacities["Residential"])
    frames.append(frame)
    utility = root / files["us_utility"]
    if utility.exists():
        frame, manifest["us_utility"] = extract_us_jedi(utility, root, "Utility", expected_capacity_mw=capacities["Utility"])
        residential = manifest["us_residential"]
        if any(manifest["us_utility"]["settings"][c] != residential["settings"][c] for c in ["B12", "B14", "B24"]):
            raise ValueError("US rooftop/utility runs have unmatched geography or year settings")
        frames.append(frame)
    else:
        manifest["us_utility"] = {"path": files["us_utility"], "status": "missing; rounded screenshots are not authoritative inputs"}
    return pd.concat(frames, ignore_index=True), manifest


def combine_curated_sources(curated: pd.DataFrame, extracted: pd.DataFrame) -> pd.DataFrame:
    """Combine manual records with workbook-owned factors without persisting copies."""
    keys = ["Source", "Technology", "Factor_Type", "Job_Type", "Year"]
    if extracted.duplicated(keys).any() or not curated.merge(extracted, on=keys).empty:
        raise ValueError("Overlapping manual and workbook factor keys")
    return pd.concat([curated, extracted], ignore_index=True)
