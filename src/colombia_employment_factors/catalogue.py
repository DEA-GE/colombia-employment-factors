"""Read approved catalogue rows once for cost ratios and technology assumptions."""
from __future__ import annotations

import csv
import math
from pathlib import Path

import openpyxl
import pandas as pd

from .mappings import (
    BASE_SOURCE_BY_TECHNOLOGY, CAPEX_OPEX_TECH_MAPPINGS,
    RUTOVITZ_2015_CONSTRUCTION_TIME_MAPPINGS, RUTOVITZ_2015_CONSTRUCTION_TIME_YEARS,
)


def _number(value) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(result):
        raise ValueError(f"Non-finite catalogue value: {value}")
    return result


def read_catalogue(path: Path, selections_path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return ratios and assumptions using saved caches and reviewed row selectors."""
    with selections_path.open(encoding="utf-8", newline="") as handle:
        selections = list(csv.DictReader(handle))
    ratios, technologies = [], {}
    with_workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        selected_sheets = [s["technology_sheet"] for s in selections]
        expected_sheets = set(with_workbook.sheetnames) - {"Guide&Cover", "Index"}
        if len(set(selected_sheets)) != len(selected_sheets) or set(selected_sheets) != expected_sheets:
            raise ValueError("Catalogue sheet coverage changed; review catalogue_rows.csv")
        for selection in selections:
            sheet = with_workbook[selection["technology_sheet"]]
            # Some source sheets have formatting out to XEV; only B:E carry inputs.
            rows = list(sheet.iter_rows(max_col=5, values_only=True))
            record = {key: selection[key] for key in ("technology_sheet", "technology_name")}
            for kind in ("capex", "opex"):
                number = int(selection[f"{kind}_row_number"])
                label, *values = rows[number - 1][1:5]
                label = str(label).strip()
                if label != selection[f"{kind}_row_used"]:
                    raise ValueError(f"Catalogue row changed: {sheet.title}!B{number}")
                record[f"{kind}_row_number"] = number
                record[f"{kind}_row_used"] = label
                costs = dict(zip((2024, 2030, 2050), map(_number, values)))
                for year, cost in costs.items():
                    record[f"{kind}_{year}"] = cost
                for year, base in ((2030, 2024), (2050, 2030)):
                    numerator, denominator = costs[year], costs[base]
                    if denominator is not None and denominator <= 0:
                        raise ValueError(f"Nonpositive catalogue denominator: {sheet.title}/{kind}/{base}")
                    record[f"{kind}_ratio_{year}_over_{base}"] = (
                        numerator / denominator if numerator is not None and denominator is not None else None
                    )
            ratios.append(record)
            metrics = {"technology_sheet": sheet.title}
            for metric, label in (("lifetime", "Technical lifetime (years)"), ("construction_time", "Construction time (years)")):
                metrics.update({f"{metric}_years": None, f"{metric}_source_year": None, f"{metric}_row_number": None})
                for number, row in enumerate(rows, 1):
                    if row[1] == label:
                        current, future = _number(row[2]), _number(row[3])
                        metrics[f"{metric}_years"] = current if current is not None else future
                        metrics[f"{metric}_source_year"] = 2024 if current is not None else 2030
                        metrics[f"{metric}_row_number"] = number
            technologies[selection["technology_name"]] = metrics
    finally:
        with_workbook.close()
    return pd.DataFrame(ratios), build_technology_assumptions(technologies)


def build_technology_assumptions(catalogue_rows: dict[str, dict]) -> pd.DataFrame:
    """Extract 2024 technology lifetimes and construction times for model technologies."""
    rows = []
    for technology in BASE_SOURCE_BY_TECHNOLOGY:
        mapped = CAPEX_OPEX_TECH_MAPPINGS.get(technology, ("", "unmapped_constant", ""))
        catalogue_tech, mapping_type, _ = mapped
        catalogue_row = catalogue_rows.get(catalogue_tech, {})
        notes = []
        lifetime_years = catalogue_row.get("lifetime_years")
        lifetime_source_year = catalogue_row.get("lifetime_source_year")
        construction_time_years = catalogue_row.get("construction_time_years")
        construction_time_source_year = catalogue_row.get("construction_time_source_year")
        source = "Colombian Technology Catalogue"

        if lifetime_source_year == 2030:
            notes.append("2024 catalogue lifetime unavailable; 2030 catalogue value used.")
        if construction_time_source_year == 2030:
            notes.append("2024 catalogue construction time unavailable; 2030 catalogue value used.")

        rutovitz_mapping = RUTOVITZ_2015_CONSTRUCTION_TIME_MAPPINGS.get(technology)
        if construction_time_years is None and rutovitz_mapping:
            construction_time_years = RUTOVITZ_2015_CONSTRUCTION_TIME_YEARS[rutovitz_mapping]
            construction_time_source_year = 2015
            source = "Rutovitz 2015 Table 1"
            notes.append("No matching catalogue row; construction time from Rutovitz 2015 Table 1.")

        if technology.startswith("Transmission"):
            notes.append(
                "No matching Technology Catalogue or Rutovitz 2015 Table 1 construction-time row."
            )

        rows.append(
            {
                "Technology": technology,
                "Catalogue_Tech": catalogue_tech,
                "lifetime_years": lifetime_years,
                "lifetime_source_year": lifetime_source_year,
                "construction_time_years": construction_time_years,
                "construction_time_source_year": construction_time_source_year,
                "construction_time_source": source if construction_time_years is not None else "",
                "technology_sheet": catalogue_row.get("technology_sheet", ""),
                "lifetime_row_number": catalogue_row.get("lifetime_row_number"),
                "construction_time_row_number": catalogue_row.get("construction_time_row_number"),
                "mapping_type": mapping_type,
                "notes": " ".join(notes),
            }
        )

    return pd.DataFrame(rows).sort_values("Technology").reset_index(drop=True)
